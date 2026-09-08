#!/usr/bin/env python3
"""
Reconstruye MyInvestor desde su Excel de movimientos.

Son dos apuntes en todo el año: una entrada de 100 € y la suscripción del fondo
Numantia por esos mismos 100 €, que deja la cuenta a cero. La suscripción NO se
registra como movimiento —en Solvento el efectivo lo mueve la operación de
inversión— pero sí se le corrige la fecha con la del extracto, que es cuando el
dinero salió de verdad.

Pequeño como es, el paso hace falta: su entrada de 100 € venía de otra cuenta, y
al reconstruir aquélla desde su extracto esta mitad se queda sin su pareja. Sin
rehacerla, MyInvestor se queda en −100 €.

    python3 generar_myinvestor.py <movimientos.xlsx> <entrada.json> <salida.json>
"""
import sys, re, json, datetime, random
import openpyxl

CUENTA = "MyInvestor"
# Conceptos que son suscripciones o reembolsos de fondos, no movimientos de caja
RE_FONDO = re.compile(r"numantia|renta 4|fondo|suscripc|reembols", re.I)


def nid(p):
    return p + "".join(random.choice("0123456789abcdef") for _ in range(10))


def main():
    xlsx, entrada, salida = sys.argv[1:4]
    wb = openpyxl.load_workbook(xlsx, data_only=True)
    sh = wb[wb.sheetnames[0]]

    filas, cabecera = [], None
    for fila in sh.iter_rows(values_only=True):
        celdas = [("" if c is None else c) for c in fila]
        txt = " ".join(str(c) for c in celdas)
        if "Fecha Operación" in txt:
            cabecera = True
            continue
        if not cabecera:
            continue
        # fecha | fecha valor | movimiento | · | importe | saldo
        vals = [c for c in celdas if str(c).strip()]
        if len(vals) < 3:
            continue
        f = str(celdas[1]).strip()
        m = re.match(r"^(\d{1,2})/(\d{1,2})/(\d{4})", f)
        if not m:
            continue
        nums = [c for c in celdas if isinstance(c, (int, float))]
        if len(nums) < 2:
            continue
        filas.append({"fecha": datetime.date(int(m.group(3)), int(m.group(2)), int(m.group(1))),
                      "concepto": str(celdas[3]).strip(), "importe": float(nums[-2]),
                      "saldo": float(nums[-1])})

    doc = json.load(open(entrada))
    # fuera lo que hubiera de esta cuenta: se rehace del extracto
    doc["movimientos"] = [m for m in doc["movimientos"]
                          if str(m.get("cuenta_origen") or "") != CUENTA
                          and str(m.get("cuenta_destino") or "") != CUENTA]
    caja, fondos = 0, 0
    for f in filas:
        if RE_FONDO.search(f["concepto"]):
            # la fecha real de la salida de dinero manda sobre la registrada
            for r in doc["inversiones"]:
                if str(r.get("cuenta") or "") != CUENTA:
                    continue
                if abs(abs(float(str(r.get("coste") or 0).replace(",", "."))) - abs(f["importe"])) < 0.005:
                    r["fecha"] = f"{f['fecha']:%d/%m/%Y}"; fondos += 1
            continue
        entra = f["importe"] > 0
        doc["movimientos"].append({
            "id": nid("m"), "marca_temporal": datetime.datetime.now().strftime("%d/%m/%Y %H:%M:%S"),
            "fecha": f"{f['fecha']:%d/%m/%Y}", "tipo": "Ingreso" if entra else "Gasto",
            "importe": f"{abs(f['importe']):.2f}",
            "cuenta_origen": "" if entra else CUENTA, "cuenta_destino": CUENTA if entra else "",
            "tipo_ingreso": "", "tipo_gasto": "", "tipo_prestamo": "", "persona_prestamo": "",
            "detalle": f["concepto"] or "Entrada de dinero en MyInvestor",
            "imp_ref": f"MyInvestor|{f['fecha']:%Y%m%d}|{f['importe']:.2f}"})
        caja += 1

    json.dump(doc, open(salida, "w"), ensure_ascii=False, indent=2)

    num = lambda x: float(str(x or 0).replace(",", "."))
    saldo = 0.0
    for m in doc["movimientos"]:
        imp = abs(num(m.get("importe")))
        o, d = str(m.get("cuenta_origen") or ""), str(m.get("cuenta_destino") or "")
        if m.get("tipo") == "Ingreso" and d == CUENTA: saldo += imp
        elif m.get("tipo") == "Gasto" and o == CUENTA: saldo -= imp
        elif m.get("tipo") == "Traspaso":
            if o == CUENTA: saldo -= imp
            if d == CUENTA: saldo += imp
    for r in doc["inversiones"]:
        if str(r.get("cuenta") or "") != CUENTA: continue
        if (r.get("tipo_movimiento") or "Compra") in ("Traspaso", "Herencia"): continue
        saldo -= num(r.get("coste"))
    saldo = round(saldo, 2)
    obj = filas[-1]["saldo"] if filas else 0.0
    print(f"{len(filas)} apuntes en el extracto · {caja} de caja · {fondos} operaciones con la fecha corregida")
    print(f"\n§2 CUADRE   extracto {obj:>9,.2f} €   reconstruido {saldo:>9,.2f} €   "
          f"{'✓ EXACTO' if abs(saldo - obj) < 0.005 else f'✗ descuadre {saldo - obj:+.2f} €'}".replace(",", " "))


if __name__ == "__main__":
    main()
